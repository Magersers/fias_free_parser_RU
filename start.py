import json
import time
import re
from pathlib import Path
from datetime import datetime, date

import requests
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from tqdm import tqdm


INPUT_FILE = "input.xlsx"
OUTPUT_FILE = "output_addresses.xlsx"
CACHE_FILE = "ahunter_cache.json"

AHUNTER_URL = "https://ahunter.ru/site/fetch/address"

# Если в первой строке заголовки — оставь 2.
# Если данные начинаются сразу с первой строки — поставь 1.
START_ROW = 2

# Пауза между запросами. Для 8000 строк лучше не ставить 0.
REQUEST_DELAY = 0.1

# Сколько раз повторять запрос при временной ошибке.
MAX_RETRIES = 3


def normalize_guid(value):
    if value is None:
        return ""

    value = str(value).strip().lower()

    match = re.search(
        r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}",
        value,
        re.IGNORECASE,
    )

    return match.group(0).lower() if match else value


def excel_date_to_string(value):
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    return str(value)


def load_cache():
    path = Path(CACHE_FILE)

    if not path.exists():
        return {}

    try:
        with path.open("r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        print("Не удалось прочитать кэш, начинаю с пустого.")
        return {}


def save_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


def ahunter_request(guid):
    params = {
        "output": "json",
        "query": f"fias:{guid}",
    }

    last_error = None

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = requests.get(
                AHUNTER_URL,
                params=params,
                timeout=30,
                headers={
                    "User-Agent": "Mozilla/5.0 FIAS address resolver"
                },
            )

            response.raise_for_status()
            return response.json()

        except Exception as e:
            last_error = e
            print(f"Ошибка запроса GUID {guid}, попытка {attempt}/{MAX_RETRIES}: {e}")
            time.sleep(1.5 * attempt)

    raise last_error


def get_field(fields, level):
    for item in fields:
        if item.get("level") == level:
            return item

    return {}


def format_field(field):
    field_type = field.get("type") or ""
    name = field.get("name") or ""

    if field_type and name:
        return f"{field_type} {name}"

    return name or ""


def extract_address_parts(raw):
    address = raw.get("address") or {}
    fields = address.get("fields") or []

    country = address.get("country") or {}
    time_zone = address.get("time_zone") or {}
    post_office = address.get("post_office") or {}

    region = get_field(fields, "Region")
    district = get_field(fields, "District")
    city = get_field(fields, "City")
    place = get_field(fields, "Place")
    site = get_field(fields, "Site")
    street = get_field(fields, "Street")
    house = get_field(fields, "House")
    building = get_field(fields, "Building")
    structure = get_field(fields, "Structure")
    flat = get_field(fields, "Flat")
    zip_code = get_field(fields, "Zip")

    return {
        "full_address": address.get("pretty") or "",

        "postal_code": zip_code.get("name") or "",

        "country_code": country.get("code") or "",
        "country_name": country.get("name") or "",
        "country_sign": country.get("sign") or "",

        "region": format_field(region),
        "region_name": region.get("name") or "",
        "region_type": region.get("type") or "",
        "region_code": region.get("c") or "",

        "district": format_field(district),
        "district_name": district.get("name") or "",
        "district_type": district.get("type") or "",
        "district_code": district.get("c") or "",

        "city": format_field(city),
        "city_name": city.get("name") or "",
        "city_type": city.get("type") or "",
        "city_code": city.get("c") or "",

        "place": format_field(place),
        "place_name": place.get("name") or "",
        "place_type": place.get("type") or "",

        "site": format_field(site),
        "site_name": site.get("name") or "",
        "site_type": site.get("type") or "",

        "street": format_field(street),
        "street_name": street.get("name") or "",
        "street_type": street.get("type") or "",
        "street_code": street.get("c") or "",

        "house": format_field(house),
        "house_name": house.get("name") or "",
        "house_type": house.get("type") or "",

        "building": format_field(building),
        "building_name": building.get("name") or "",
        "building_type": building.get("type") or "",

        "structure": format_field(structure),
        "structure_name": structure.get("name") or "",
        "structure_type": structure.get("type") or "",

        "flat": format_field(flat),
        "flat_name": flat.get("name") or "",
        "flat_type": flat.get("type") or "",

        "timezone": time_zone.get("name") or "",
        "timezone_utc": time_zone.get("utc_zone") or "",
        "timezone_msk": time_zone.get("msk_zone") or "",

        "post_office_address": post_office.get("pretty") or "",
        "post_office_lat": post_office.get("lat") or "",
        "post_office_lon": post_office.get("lon") or "",
    }


def read_input_rows(filename):
    wb = load_workbook(filename, data_only=True)
    ws = wb.active

    rows = []

    for row_num in range(START_ROW, ws.max_row + 1):
        guid = normalize_guid(ws.cell(row=row_num, column=1).value)
        date_value = excel_date_to_string(ws.cell(row=row_num, column=2).value)

        if not guid:
            continue

        rows.append({
            "source_row": row_num,
            "guid": guid,
            "date": date_value,
        })

    return rows


def get_headers():
    """
    input_fias — предпоследняя колонка.
    date — последняя колонка.
    """
    return [
        "source_row",

        "full_address",
        "postal_code",

        "country_code",
        "country_name",
        "country_sign",

        "region",
        "region_name",
        "region_type",
        "region_code",

        "district",
        "district_name",
        "district_type",
        "district_code",

        "city",
        "city_name",
        "city_type",
        "city_code",

        "place",
        "place_name",
        "place_type",

        "site",
        "site_name",
        "site_type",

        "street",
        "street_name",
        "street_type",
        "street_code",

        "house",
        "house_name",
        "house_type",

        "building",
        "building_name",
        "building_type",

        "structure",
        "structure_name",
        "structure_type",

        "flat",
        "flat_name",
        "flat_type",

        "timezone",
        "timezone_utc",
        "timezone_msk",

        "post_office_address",
        "post_office_lat",
        "post_office_lon",

        "status",
        "error",
        "raw_json",

        "input_fias",
        "date",
    ]


def write_output(rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "addresses"

    headers = get_headers()

    ws.append(headers)

    for item in rows:
        ws.append([item.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="D9EAF7")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)

        if header in ("full_address", "post_office_address", "raw_json", "error"):
            ws.column_dimensions[col_letter].width = 60
        elif header == "input_fias":
            ws.column_dimensions[col_letter].width = 40
        elif header == "date":
            ws.column_dimensions[col_letter].width = 18
        elif header.endswith("_name"):
            ws.column_dimensions[col_letter].width = 26
        else:
            ws.column_dimensions[col_letter].width = 18

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(filename)


def main():
    input_rows = read_input_rows(INPUT_FILE)

    print(f"Найдено строк для обработки: {len(input_rows)}")

    cache = load_cache()
    output_rows = []

    for row in tqdm(input_rows, desc="Обработка ФИАС"):
        guid = row["guid"]

        try:
            if guid in cache:
                raw = cache[guid]
            else:
                raw = ahunter_request(guid)
                cache[guid] = raw
                save_cache(cache)
                time.sleep(REQUEST_DELAY)

            parts = extract_address_parts(raw)

            if not parts.get("full_address"):
                status = "not_found"
                error = "AHunter не вернул address.pretty"
            else:
                status = "ok"
                error = ""

            output_rows.append({
                "source_row": row["source_row"],
                **parts,
                "status": status,
                "error": error,
                "raw_json": json.dumps(raw, ensure_ascii=False),
                "input_fias": guid,
                "date": row["date"],
            })

        except Exception as e:
            output_rows.append({
                "source_row": row["source_row"],
                "status": "error",
                "error": str(e),
                "raw_json": "",
                "input_fias": guid,
                "date": row["date"],
            })

            time.sleep(REQUEST_DELAY)

    write_output(output_rows, OUTPUT_FILE)

    print("")
    print(f"Готово: {OUTPUT_FILE}")
    print(f"Обработано строк: {len(output_rows)}")
    print(f"Кэш сохранён: {CACHE_FILE}")


if __name__ == "__main__":
    main()